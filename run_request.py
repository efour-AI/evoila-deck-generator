"""Build one walking deck for an Airtable 'Deck Requests' record.

Run by GitHub Actions (see .github/workflows/build-deck.yml).
Env: AIRTABLE_TOKEN (secret), RECORD_ID, optional RUN_URL, MODE (build-deck or check-library).

MODE=check-library: RECORD_ID is a 'Slide Library Versions' row. Builds the full library from that
Word doc, writes a check report, attaches the library PowerPoint, and makes the version Live
(retiring the previous Live one) unless the doc could not be read.

1. Read the request row and its linked presenters from Airtable.
2. Download the master template and the engine zip (slide builders + icons) from the 'Deck Assets' table.
3. Build the PPTX with engine/build_library.py, add AI talking points to speaker notes.
4. Convert to PDF with LibreOffice.
5. Upload both files back to the row and set Status = Ready (or Error with detail).
"""
import base64, datetime, io, json, os, re, subprocess, sys, traceback, zipfile
import urllib.request, urllib.parse

BASE = "appYVahdfC3gbtArV"
T_REQ = "tbl3bFM0jwCjANMYO"
T_PRES = "tblWQmz1MD437ybqk"
T_ASSETS = "tblVXLTTiAInLHs0E"
T_LIB = "tblQTgldqzt69bQrf"     # Slide Library Versions
L = {"version": "fldmi2UVTliwrN4SQ", "doc": "flduUUR0au4ra3BYw", "status": "flda45mlCGY34rDdP",
     "report": "fldIcYWdhTZmRpThL", "pptx": "fld6KzeSNicCGChq3", "log": "fldYwqeziIcjaTnE3", "added": "fld1eYcvwCTDRcM57"}
F = {  # Deck Requests field IDs
    "name": "fldXM03WhdelyP62D", "email": "fldsh54IoNqdFDRM4", "prospect": "fldh6PvIYkHgR7m8o",
    "date": "fldehcJBsAzui66pz", "audience": "fldG2ZAIhyN5KdFAy", "framing": "fldCEoEULzI0Px2BR",
    "slides": "fldMOj09vHxxvahs5", "notes": "fldtrr3cxVa29DnrF", "status": "fldxQZCQ8OHA7sOTa",
    "plan": "fldE9rNtby6JwC4qk", "pptx": "fldxwsdZxpMOq0NiL", "pdf": "fldNi5sd61xg5OhcL",
    "log": "fldtBHsDTgcLuNQhb", "error": "fldELZ4iyycJu7mOD", "presenters": "fldyMIZMHeWaET7sz",
    "libver": "fldxnMaaGGwVKqcb4", "website": "fldELJzcC0UzkREuI", "logo": "fldcLJ4rbL8RQ3NGE",
    "research": "fldgjfqb7yBy9ziyq", "rcheck": "fldqFZN5K6Xgi0A5i",
}
P = {"name": "fldr4qcJSOTaspQGm", "title": "fld68D6bAxXqefbSR", "email": "fldv8VzsSsoySdCds"}
A = {"name": "fldwyRdfc2x7SRlat", "file": "fld3T1O2M6dsaCeSy", "code": "fldNtFG2NXXI44B69"}

TOKEN = os.environ["AIRTABLE_TOKEN"]
REC = os.environ["RECORD_ID"].strip()
RUN_URL = os.environ.get("RUN_URL", "")
MODE = os.environ.get("MODE", "build-deck") or "build-deck"
HERE = os.path.dirname(os.path.abspath(__file__))
WORK = os.path.join(HERE, "work")
ENGINE = os.path.join(HERE, "engine")


def api(method, url, body=None):
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(url, data=data, method=method, headers={
        "Authorization": f"Bearer {TOKEN}", "Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=120) as r:
        return json.loads(r.read() or b"{}")


def get_record(table, rid):
    return api("GET", f"https://api.airtable.com/v0/{BASE}/{table}/{rid}?returnFieldsByFieldId=true")["fields"]


def update(fields, table=None, rid=None):
    api("PATCH", f"https://api.airtable.com/v0/{BASE}/{table or T_REQ}/{rid or REC}",
        {"fields": fields, "typecast": True})


def upload(field, path, ctype, rid=None):
    with open(path, "rb") as fh:
        b64 = base64.b64encode(fh.read()).decode()
    api("POST", f"https://content.airtable.com/v0/{BASE}/{rid or REC}/{field}/uploadAttachment",
        {"contentType": ctype, "file": b64, "filename": os.path.basename(path)})


def download(url, path):
    with urllib.request.urlopen(url, timeout=300) as r, open(path, "wb") as fh:
        fh.write(r.read())


def fetch_assets():
    rows = api("GET", f"https://api.airtable.com/v0/{BASE}/{T_ASSETS}?returnFieldsByFieldId=true")["records"]
    by = {r["fields"].get(A["name"], ""): r["fields"].get(A["file"], []) for r in rows}
    tpl = os.path.join(WORK, "template.pptx")
    download(by["Master template"][0]["url"], tpl)
    z = os.path.join(WORK, "assets.zip")
    download(by["Engine and icons"][0]["url"], z)
    zipfile.ZipFile(z).extractall(ENGINE)   # build_library.py, pptkit.py, icons/, earth_src.jpeg
    for r in rows:                           # newer engine code kept as text on the row wins over the zip copy
        code = (r["fields"].get(A["code"]) or "").strip()
        if code:
            zipfile.ZipFile(io.BytesIO(base64.b64decode(code))).extractall(ENGINE)
    return tpl


def live_library():
    """Newest Live row of Slide Library Versions -> (local docx path, version) or (None, None)."""
    rows = api("GET", f"https://api.airtable.com/v0/{BASE}/{T_LIB}?returnFieldsByFieldId=true")["records"]
    live = [r for r in rows if r["fields"].get(L["status"]) == "Live" and r["fields"].get(L["doc"])]
    if not live:
        return None, None
    live.sort(key=lambda r: r["fields"].get(L["added"]) or r.get("createdTime", ""), reverse=True)
    f = live[0]["fields"]
    path = os.path.join(WORK, "library.docx")
    download(f[L["doc"]][0]["url"], path)
    return path, (f.get(L["version"]) or "").strip() or live[0]["id"]


def report_text(r):
    out = []
    if r["errors"]:
        out += ["PROBLEMS (this version was not made live)"] + [f"- {x}" for x in r["errors"]] + [""]
    out += ["MISSING LINES (the deck used the previous wording)"] + ([f"- {x}" for x in r["fallbacks"]] or ["- None"]) + [""]
    out += ["STILL IN BRACKETS (will print as written until an owner fills them)"] + ([f"- {x}" for x in r["placeholders"]] or ["- None"]) + [""]
    out += ["LINES THE GENERATOR DOES NOT USE (a new label, a typo in a label, or a design note)"] + ([f"- {x}" for x in r["unused"]] or ["- None"])
    return "\n".join(out)


def check_library():
    """MODE=check-library: build the full library from one version row and report."""
    os.makedirs(WORK, exist_ok=True)
    f = get_record(T_LIB, REC)
    files = f.get(L["doc"]) or []
    if not files or not files[0].get("filename", "").lower().endswith(".docx"):
        update({L["status"]: "Problems found", L["log"]: RUN_URL or None,
                L["report"]: "PROBLEMS (this version was not made live)\n- Attach the Word file (.docx), not a PDF or other format."}, T_LIB)
        print("check: no docx")
        return
    doc = os.path.join(WORK, "library.docx")
    download(files[0]["url"], doc)
    version = (f.get(L["version"]) or "").strip() or "this version"
    tpl = fetch_assets()
    out = os.path.join(WORK, f"evoila_Walking_Deck_Library_{safe(version)}.pptx")
    rep_path = os.path.join(WORK, "report.json")
    try:
        subprocess.run([sys.executable, os.path.join(ENGINE, "build_library.py"), tpl, out, "--library", doc,
                        "--version", version, "--report", rep_path], check=True,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
        r = json.load(open(rep_path))
    except Exception:
        r = {"errors": ["The Word file could not be read as a slide library. Check that each slide still has its table "
                        "(Variant | On-slide copy | Design spec | Speaker notes)."], "fallbacks": [], "placeholders": [], "unused": []}
    ok = not r["errors"]
    update({L["pptx"]: []}, T_LIB)
    if ok and os.path.exists(out):
        upload(L["pptx"], out, "application/vnd.openxmlformats-officedocument.presentationml.presentation")
        # retire the previous Live version(s)
        rows = api("GET", f"https://api.airtable.com/v0/{BASE}/{T_LIB}?returnFieldsByFieldId=true")["records"]
        for row in rows:
            if row["id"] != REC and row["fields"].get(L["status"]) == "Live":
                update({L["status"]: "Retired"}, T_LIB, row["id"])
    update({L["status"]: "Live" if ok else "Problems found", L["report"]: report_text(r)[:90000],
            L["log"]: RUN_URL or None}, T_LIB)
    print("check:", "live" if ok else "problems")   # no doc content in logs; the repo is public


UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36"}


def fetch(url, timeout=20):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=timeout) as r:
        return r.read(), r.headers.get("Content-Type", ""), r.geturl()


def to_png(data, ctype, url, out):
    """Save image bytes as PNG (SVG converted). Returns path or None if it is not a usable image."""
    from PIL import Image
    if "svg" in ctype or url.lower().split("?")[0].endswith(".svg") or data.lstrip()[:5] in (b"<svg ", b"<?xml"):
        import cairosvg
        cairosvg.svg2png(bytestring=data, write_to=out, output_width=800)
    else:
        with open(out, "wb") as fh:
            fh.write(data)
    im = Image.open(out)
    if max(im.size) < 120:          # favicons and tiny icons look bad on a slide
        return None
    thumb = im.convert("RGB").resize((64, 64))
    if len(set(thumb.getdata())) > 1800:   # photo, not a logo
        return None
    im.convert("RGBA").save(out)
    return out


def find_logo(site, out):
    """Best effort: the company's logo from its home page. Returns a PNG path or None."""
    import html as _h
    from urllib.parse import urljoin, urlparse
    if not re.match(r"https?://", site):
        site = "https://" + site.strip().strip("/")
    try:
        page, _, base = fetch(site)
        doc = page.decode("utf-8", "ignore")
    except Exception:                            # site blocks automated visits: use its site icon from Google
        try:
            data, ctype, real = fetch("https://www.google.com/s2/favicons?sz=256&domain=" + urlparse(site).netloc)
            return to_png(data, ctype, real, out)
        except Exception:
            return None
    cands = []                                   # (kind, value): kind 'url' or 'svg'

    def img_src(tag):
        m = (re.search(r"""\bsrc=["']([^"']+)["']""", tag, re.I) or re.search(r"""data-src=["']([^"']+)["']""", tag, re.I)
             or re.search(r"""srcset=["']([^"'\s,]+)""", tag, re.I))
        if not m or m.group(1).startswith("data:") or re.search(r"\.jpe?g(\?|$)", m.group(1), re.I):
            return None
        return m.group(1)

    # 1. <img> whose file, class or id says logo/brand
    for tag in re.findall(r"<img\b[^>]*>", doc, re.I):
        src = img_src(tag)
        attrs = " ".join(re.findall(r"""\b(?:class|id)=["']([^"']*)["']""", tag, re.I))
        if src and (re.search(r"logo|brand", src, re.I) or re.search(r"logo|brand", attrs, re.I)):
            cands.append(("url", src))
    # 2. first image or inline svg inside an element labelled logo/brand (e.g. <a class="header__brand">)
    for m in re.finditer(r"""\b(?:class|id|aria-label)=["'][^"']*(?:logo|brand)[^"']*["'][^>]*>""", doc, re.I):
        win = doc[m.end(): m.end() + 2500]
        im = re.search(r"<img\b[^>]*>", win, re.I)
        sv = re.search(r"<svg\b.*?</svg>", win, re.I | re.S)
        if sv and (not im or sv.start() < im.start()):
            cands.append(("svg", sv.group(0)))
        elif im and img_src(im.group(0)):
            cands.append(("url", img_src(im.group(0))))
    # 3. first image inside <header>
    hm = re.search(r"<header\b.*?</header>", doc, re.I | re.S)
    if hm:
        im = re.search(r"<img\b[^>]*>", hm.group(0), re.I)
        if im and img_src(im.group(0)):
            cands.append(("url", img_src(im.group(0))))
    # 4. app icons
    for m in re.finditer(r"""<link\b[^>]*rel=["'][^"']*(?:apple-touch-icon|icon)[^"']*["'][^>]*>""", doc, re.I):
        h = re.search(r"""href=["']([^"']+)["']""", m.group(0), re.I)
        if h:
            cands.append(("url", h.group(1)))
    cands.append(("url", "/apple-touch-icon.png"))
    cands.append(("url", "https://www.google.com/s2/favicons?sz=256&domain=" + urlparse(base).netloc))
    seen = set()
    for kind, v in cands:
        if v in seen:
            continue
        seen.add(v)
        if len(seen) > 12:
            break
        try:
            if kind == "svg":
                svg = v if "xmlns" in v else v.replace("<svg", '<svg xmlns="http://www.w3.org/2000/svg"', 1)
                p = to_png(svg.encode(), "image/svg+xml", "inline.svg", out)
            else:
                data, ctype, real = fetch(urljoin(base, _h.unescape(v)))
                p = to_png(data, ctype, real, out)
            if p:
                return p
        except Exception:
            continue
    return None


def company_logo(f):
    """Logo file on the row wins; otherwise look on the company website. Never fails the build."""
    out = os.path.join(WORK, "company_logo.png")
    try:
        files = f.get(F["logo"]) or []
        if files:
            data, ctype, real = fetch(files[0]["url"], timeout=60)
            return to_png(data, files[0].get("type", "") or ctype, files[0].get("filename", real), out)
        if f.get(F["website"]):
            p = find_logo(f[F["website"]], out)
            if p:                          # keep it on the row so people can see it and swap it
                upload(F["logo"], p, "image/png")
            return p
    except Exception:
        return None
    return None


PAGE_DATE = {}      # url -> publication date found on the page itself (or None)
MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def parse_date(v):
    """'2026-03-15', '2026-03', '2026-03-15T10:00Z', 'March 2026', 'Mar 15, 2026' -> date, or None if no month is given."""
    v = str(v or "").strip()
    m = re.match(r"(\d{4})-(\d{1,2})(?:-(\d{1,2}))?", v)
    if m:
        try:
            return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3) or 1))
        except ValueError:
            return None
    m = re.search(r"\b([A-Za-z]{3})[a-z]*\.?\s+(?:(\d{1,2}),?\s+)?(\d{4})\b", v)
    if m and m.group(1).lower() in MONTHS:
        try:
            return datetime.date(int(m.group(3)), MONTHS[m.group(1).lower()], int(m.group(2) or 1))
        except ValueError:
            return None
    return None


DATE_PATTERNS = [
    r"<meta[^>]+(?:property|name|itemprop)=[\"'](?:article:published_time|datePublished|citation_publication_date|citation_date|dc\.date(?:\.issued)?|pubdate|publish[-_]?date|date)[\"'][^>]*content=[\"']([^\"']+)[\"']",
    r"<meta[^>]+content=[\"']([^\"']+)[\"'][^>]*(?:property|name|itemprop)=[\"'](?:article:published_time|datePublished|citation_publication_date|pubdate|publish[-_]?date)[\"']",
    r"\"datePublished\"\s*:\s*\"([^\"]+)\"",
    r"<time[^>]+datetime=[\"']([^\"']+)[\"']",
]


def html_published(doc):
    """Publication date the page declares about itself (meta tags, JSON-LD datePublished, <time>)."""
    for pat in DATE_PATTERNS:
        m = re.search(pat, doc, re.I)
        if m and parse_date(m.group(1)):
            return parse_date(m.group(1))
    return None


BOT_WALL = re.compile(r"just a moment|verify you are human|performing security verification|enable javascript and cookies|access denied|captcha", re.I)


def page_text(url):
    """Visible text of a web page or PDF, lowercased. Sites that block automated visits are read through
    the r.jina.ai reader instead; a page still behind a bot wall returns '' (the figure is then dropped)."""
    try:
        txt = page_text_direct(url)
        if len(txt) > 1500 and not BOT_WALL.search(txt[:3000]):
            return txt
    except Exception:
        pass
    data, _, _ = fetch("https://r.jina.ai/" + url, timeout=60)
    doc = data.decode("utf-8", "ignore")
    head, _, body = doc.partition("Markdown Content:")
    if "requiring CAPTCHA" in head or BOT_WALL.search(body[:1500]):
        PAGE_DATE[url] = None
        return ""
    m = re.search(r"Published Time:\s*(\S+)", head)
    PAGE_DATE[url] = parse_date(m.group(1)) if m else None
    return re.sub(r"\s+", " ", body).lower()


def page_text_direct(url):
    """Visible text of a web page or PDF fetched directly, lowercased."""
    data, ctype, real = fetch(url, timeout=30)
    if "pdf" in ctype or real.lower().split("?")[0].endswith(".pdf") or data[:4] == b"%PDF":
        import pypdf
        rd = pypdf.PdfReader(io.BytesIO(data))
        txt = " ".join((pg.extract_text() or "") for pg in rd.pages[:60])
        try:
            cd = rd.metadata.creation_date if rd.metadata else None
            PAGE_DATE[url] = cd.date() if cd else None
        except Exception:
            PAGE_DATE[url] = None
    else:
        txt = data.decode("utf-8", "ignore")
        PAGE_DATE[url] = html_published(txt)
        txt = re.sub(r"<script.*?</script>|<style.*?</style>", " ", txt, flags=re.S | re.I)
        txt = re.sub(r"<[^>]+>", " ", txt)
        import html as _h
        txt = _h.unescape(txt)
    return re.sub(r"\s+", " ", txt).lower()


def figure_core(fig):
    """'$4.88 million' -> ['4.88'], '1 in 3' -> ['1 in 3'], '48%' -> ['48']."""
    f = (fig or "").lower().replace(",", "")
    m = re.search(r"\d+\s+(?:in|of|out of)\s+\d+", f)
    if m:
        return [m.group(0)]
    return re.findall(r"\d+(?:\.\d+)?", f)[:1]


STOP = set("about after their there which would could should these those with from that this have been were into more than most also only over such they them what when where your ours".split())


def figure_on_page(body, cores, claim):
    """The figure must appear on the page near at least two words of its claim (or one, for short claims)."""
    words = [w for w in re.findall(r"[a-z]{5,}", (claim or "").lower()) if w not in STOP]
    need = 2 if len(words) >= 4 else 1
    for core in cores:
        for m in re.finditer(r"(?<![\d.])" + re.escape(core) + r"(?![\d])", body):
            win = body[max(0, m.start() - 250): m.end() + 250]
            if sum(1 for w in set(words) if w[:6] in win) >= need:
                return True
    return False


def compact_figure(fig):
    """'$7.42 million' -> '$7.42M', 'USD 4.4 billion' -> '$4.4B' so the number fits its tile on one line."""
    f = re.sub(r"^USD\s*", "$", str(fig).strip())
    f = re.sub(r"\s*(million|mn)\b", "M", f, flags=re.I)
    f = re.sub(r"\s*(billion|bn)\b", "B", f, flags=re.I)
    f = re.sub(r"\s*(trillion)\b", "T", f, flags=re.I)
    return f


def drop_lead_figure(claim, fig, cores):
    """The big number already sits above the claim: drop it if the claim starts with it."""
    c = str(claim).strip()
    for core in sorted(set([fig] + list(cores or [])), key=len, reverse=True):
        if core and c.lower().startswith(str(core).lower()):
            c = c[len(core):].lstrip(" ,:-")
            c = re.sub(r"^(million|billion|M|B)\b\s*", "", c)
            break
    return c[:1].upper() + c[1:] if c else str(claim)


def verify_market(raw):
    """Keep only researched figures whose number appears on the cited page. Returns (market or None, report)."""
    try:
        m = json.loads(re.search(r"\{.*\}", raw or "", re.S).group(0))
    except Exception:
        return None, "No usable research on the row."
    if not m.get("relevant", True):
        return None, "Research step judged a data slide not useful for this meeting."
    kept, rep = [], []
    today = datetime.date.today()
    for st in (m.get("stats") or [])[:6]:
        url, cores = (st.get("url") or "").strip(), figure_core(st.get("figure"))
        label = f"{st.get('figure','')} {st.get('claim','')[:80]} ({st.get('source','')})"
        if not url.startswith("http") or not cores:
            rep.append(f"DROPPED (no source link or no number): {label}")
            continue
        if re.search(r"/(solutions?|products?|platform|pricing|demo|lp)(/|$)", urllib.parse.urlparse(url).path.lower()):
            rep.append(f"DROPPED (vendor product page, not a report): {label} {url}")
            continue
        host = urllib.parse.urlparse(url).netloc.lower()
        if re.search(r"/(blog|blogs|resources|insights)/", urllib.parse.urlparse(url).path.lower()) and not re.search(r"\.(gov|edu|org)$", host):
            rep.append(f"DROPPED (company blog, not an independent source): {label} {url}")
            continue
        try:
            body = page_text(url).replace(",", "")
            ok = figure_on_page(body, cores, st.get("claim", ""))
        except Exception:
            ok = False
        # data must be less than 12 months old: the publication date on the page itself, else the date the research gave;
        # with both, the older one counts. No month-level date at all -> dropped.
        dates = [x for x in [PAGE_DATE.get(url), parse_date(st.get("published"))] if x and x <= today + datetime.timedelta(days=31)]
        pub = min(dates) if dates else None
        if ok and not pub:
            rep.append(f"DROPPED (no publication date could be found): {label} {url}")
            continue
        if ok and (today - pub).days > 365:
            rep.append(f"DROPPED (published {pub.strftime('%b %Y')}, older than 12 months): {label} {url}")
            continue
        if ok:
            st["year"] = pub.strftime("%b %Y")      # printed after the source, e.g. "Source: HIPAA Journal, Mar 2026"
            st["figure"] = compact_figure(st.get("figure", ""))
            st["claim"] = drop_lead_figure(st.get("claim", ""), st["figure"], cores)
            kept.append(st)
            rep.append(f"USED (found on the page, published {pub.strftime('%b %Y')}): {label} {url}")
        else:
            rep.append(f"DROPPED (could not confirm on the page): {label} {url}")
        if len(kept) == 3:
            break
    if len(kept) < 2:
        return None, "\n".join(rep + ["No data slide: fewer than two figures could be confirmed."])
    m["stats"] = kept
    for k in ("headline", "subhead", "takeaway", "talking_point"):     # house rule: no em dashes on slides
        if isinstance(m.get(k), str):
            m[k] = re.sub(r"\s*[\u2014\u2013]\s*", ", ", m[k])
    for st in kept:
        st["claim"] = re.sub(r"\s*[\u2014\u2013]\s*", ", ", st.get("claim", ""))
    return m, "\n".join(rep)


def clean_attendee(a):
    """'Mickey Lettuce-CIO' -> 'Mickey Lettuce, CIO'. Hyphenated surnames are left alone."""
    a = re.sub(r"\s+", " ", str(a)).strip(" ,-")
    m = re.match(r"^(.*?[a-z])\s*[-:/]\s*([A-Z]{2,5}\b.*|(?:Chief|VP|Vice|Director|Head|Senior|Manager)\b.*)$", a)
    if m and "," not in a:
        a = f"{m.group(1).strip()}, {m.group(2).strip()}"
    return a


RESEARCH_WAIT = 360     # seconds the build waits for the research step, counted from the start of this script
T0 = datetime.datetime.now()


def wait_for_research():
    """The research step runs in Airtable while this build starts. Poll the row until it writes the research
    (or the wait runs out, in which case the deck is built without a data slide)."""
    import time
    while True:
        raw = (get_record(T_REQ, REC).get(F["research"]) or "").strip()
        if raw:
            return raw
        if (datetime.datetime.now() - T0).total_seconds() > RESEARCH_WAIT:
            return ""
        time.sleep(15)


def slide_no(label):
    m = re.match(r"\s*(\d+)", label or "")
    return int(m.group(1)) if m else None


def parse_plan(raw):
    """AI plan JSON written by the Airtable automation. Tolerates code fences and junk."""
    if not raw:
        return {}
    raw = raw.strip()
    m = re.search(r"\{.*\}", raw, re.S)
    try:
        return json.loads(m.group(0) if m else raw)
    except Exception:
        return {}


def safe(s):
    return re.sub(r"[^A-Za-z0-9]+", "_", s or "").strip("_")[:40] or "Meeting"


def main():
    os.makedirs(WORK, exist_ok=True)
    f = get_record(T_REQ, REC)
    audience = (f.get(F["audience"]) or "Executive")
    framing = (f.get(F["framing"]) or "Company")
    prospect = (f.get(F["prospect"]) or "").strip()
    date_s = f.get(F["date"])
    plan = parse_plan(f.get(F["plan"]))

    # slide order: explicit form choice wins, then AI plan, then the full 13
    chosen = sorted(n for n in (slide_no(s) for s in f.get(F["slides"], [])) if n)
    ai_slides = [slide_no(str(s.get("slide"))) for s in plan.get("slides", []) if isinstance(s, dict)]
    ai_slides = [n for n in ai_slides if n and 1 <= n <= 13]
    slides = chosen or list(dict.fromkeys(ai_slides)) or list(range(1, 14))
    if 1 not in slides:
        slides = [1] + slides           # every deck opens on the cover
    if 13 not in slides:
        slides = slides + [13]          # and closes on How to engage
    talking = {}
    for s in plan.get("slides", []):
        if isinstance(s, dict) and slide_no(str(s.get("slide"))):
            pts = s.get("talking_points") or []
            talking[slide_no(str(s.get("slide")))] = pts if isinstance(pts, list) else [str(pts)]

    # presenters and contacts
    people = []
    for pid in f.get(F["presenters"], [])[:2]:
        p = get_record(T_PRES, pid)
        people.append({"name": p.get(P["name"], ""), "title": p.get(P["title"], ""), "email": p.get(P["email"], "")})
    when = ""
    if date_s:
        d = datetime.date.fromisoformat(date_s)
        when = f"{d.strftime('%B')} {d.day}, {d.year}"
    meeting_line = "  |  ".join(x for x in [f"Prepared for {prospect}" if prospect else "", when] if x) or "[Meeting line]"
    req = {"audience": audience, "framing": framing, "slides": slides, "meeting_line": meeting_line,
           "prospect": prospect or None, "when": when,
           "presenters": "  |  ".join(f"{p['name']}, {p['title']}" for p in people) or "[Presenter name, title]"}
    if people:
        req["contacts"] = people
    logo = company_logo(f)
    if logo:
        req["logo"] = logo
    att = plan.get("attendees") or []
    req["attendees"] = [clean_attendee(a) for a in att if str(a).strip()][:3] if isinstance(att, list) else []
    market, mreport = (None, "")
    research = wait_for_research()
    if research:
        try:
            market, mreport = verify_market(research)
        except Exception:
            market, mreport = None, "The data check failed; no data slide."
    else:
        mreport = "The research step had not finished after six minutes (or failed); the deck was built without a data slide."
    if market:
        req["market"] = market
        slides = [1, "M"] + [n for n in slides if n != 1]
        req["slides"] = slides
    req_path = os.path.join(WORK, "request.json")
    json.dump(req, open(req_path, "w"))

    tpl = fetch_assets()
    lib_doc, lib_ver = live_library()
    lib_args = ["--library", lib_doc, "--version", lib_ver] if lib_doc else []
    stamp = (date_s or datetime.date.today().isoformat())
    base_name = f"evoila_{safe(prospect)}_{audience.replace(' ', '')}_{stamp}"
    pptx = os.path.join(WORK, base_name + ".pptx")
    subprocess.run([sys.executable, os.path.join(ENGINE, "build_library.py"), tpl, pptx, "--request", req_path] + lib_args, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

    # meeting talking points go on top of the library notes
    from pptx import Presentation
    prs = Presentation(pptx)
    for slide, n in zip(prs.slides, slides):     # "M" (data slide) carries its own notes
        pts = [p for p in talking.get(n, []) if str(p).strip()]
        if not pts:
            continue
        tf = slide.notes_slide.notes_text_frame
        lib = tf.text
        tf.text = ("FOR THIS MEETING\n" + "\n".join(f"- {p}" for p in pts)
                   + ("\n\nLIBRARY NOTES\n" + lib if lib.strip() else ""))
    if plan.get("summary"):
        tf = prs.slides[0].notes_slide.notes_text_frame
        tf.text = "DECK PLAN\n" + plan["summary"] + "\n\n" + tf.text
    prs.save(pptx)

    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", WORK, pptx],
                   check=True, timeout=240, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf = pptx[:-5] + ".pdf"

    update({F["pptx"]: [], F["pdf"]: []})   # a rebuild replaces the old files instead of adding to them
    upload(F["pptx"], pptx, "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    upload(F["pdf"], pdf, "application/pdf")
    update({F["status"]: "Ready", F["log"]: RUN_URL or None, F["error"]: "",
            F["libver"]: lib_ver or "built-in wording (no Live library version)",
            F["rcheck"]: mreport or None})
    print("done:", len(slides), "slides")   # no company names in logs; the repo is public


def mark_failed():
    """Called by the workflow if the job died before the script could record its own result."""
    if MODE == "check-library":
        f = get_record(T_LIB, REC)
        if f.get(L["status"]) == "Checking":
            update({L["status"]: "Problems found", L["report"]: "The check server failed or timed out. Clear Status on this row to run the check again."}, T_LIB)
    else:
        f = get_record(T_REQ, REC)
        if f.get(F["status"]) == "Building":
            update({F["status"]: "Error", F["error"]: "The build server failed or timed out. Clear Status on this row to try again."})


if __name__ == "__main__":
    if "--mark-failed" in sys.argv:
        mark_failed()
        sys.exit(0)
    try:
        check_library() if MODE == "check-library" else main()
    except Exception:
        err = traceback.format_exc()
        print("Run failed. Details are on the Airtable row.")   # keep logs free of request data
        try:
            if MODE == "check-library":
                update({L["status"]: "Problems found", L["report"]: "The check crashed:\n" + err[-3000:], L["log"]: RUN_URL or None}, T_LIB)
            else:
                update({F["status"]: "Error", F["error"]: err[-3000:], F["log"]: RUN_URL or None})
        except Exception:
            print("Could not write the error to Airtable.")
        sys.exit(1)
