"""Custom outline decks: the requester writes the outline, Claude writes the slide copy, we lay it out on brand.

Fired for Deck Requests rows with Request type = "Custom outline". The hub's deck outline page sends a JSON
outline (deck title plus one entry per slide: a layout, or "auto", and what the requester wants on it).

1. Claude turns the outline into slide copy that follows the evoila rules (make_spec).
2. render() draws each slide on the evoila master template with the same helpers as the walking deck
   (pptkit.py and the template plumbing in build_library.py, both from the engine zip on Deck Assets).
3. run() converts to PDF, uploads both files to the row and sets Status = Ready; the existing
   "Deck ready: email requester" automation sends them.

Research slides: an outline entry with layout "research" carries a topic and the requester's rules (how recent,
where it applies, which kinds of source, anything else). Claude researches it with web search, every figure is
checked on its source page with the same rules as the walking deck's data slides (trusted sources, the figure
must appear on the page, published and covering data inside the requester's window), and up to four confirmed
figures go on a Big numbers slide with their sources. With fewer than two confirmed figures the slide keeps
[figure] placeholders and the notes say why; the full check report goes to the row's Market data check field.

Nothing here prints request data: the repo is public and the build logs are visible.
"""
import datetime, json, math, os, re, subprocess, sys

import run_request as rr
import ai_steps

F_TYPE = "fldQahyDEnFyMIm87"
F_OUTLINE = "fldokfA560RE3GLh4"
F_ROOM = "fldh3Ppbw0ItJoTFw"
F_NEXT = "fldFr9808gsFms3qY"
F_BACKGROUND = "fldiUPLZFpHCY7ArN"

LAYOUTS = ["cover", "agenda", "statement", "bullets", "columns", "cards", "numbers", "steps", "comparison",
           "quote", "closing"]
ICON_NAMES = ["ArrowRight", "Award", "BadgeCheck", "Blocks", "Boxes", "Building2", "CalendarCheck", "ChartLine",
              "CircleCheck", "ClipboardCheck", "Cloud", "Coins", "Container", "Cpu", "Database", "Factory", "Gauge",
              "GitBranch", "Globe", "GraduationCap", "Handshake", "HardDrive", "HeartPulse", "Landmark", "Laptop",
              "Layers", "LifeBuoy", "Lock", "Mail", "MapPin", "MessagesSquare", "Monitor", "Network", "PackageCheck",
              "PiggyBank", "Recycle", "RefreshCw", "Rocket", "Route", "Scaling", "Search", "Server", "ServerCog",
              "Settings", "Shield", "ShieldCheck", "Sparkles", "Target", "Timer", "TrendingDown", "TrendingUp",
              "TriangleAlert", "UserCheck", "Users", "Workflow", "Zap"]

# ---------------------------------------------------------------- Claude: outline -> slide copy

ITEM = {"type": "object", "properties": {
    "heading": {"type": "string"}, "text": {"type": "string"},
    "points": {"type": "array", "items": {"type": "string"}},
    "figure": {"type": "string"}, "icon": {"type": "string", "enum": ICON_NAMES}}}
SPEC_SCHEMA = {
    "type": "object",
    "properties": {
        "title": {"type": "string"}, "subtitle": {"type": "string"},
        "slides": {"type": "array", "items": {"type": "object", "properties": {
            "layout": {"type": "string", "enum": LAYOUTS},
            "title": {"type": "string"}, "kicker": {"type": "string"},
            "statement": {"type": "string"}, "support": {"type": "string"},
            "items": {"type": "array", "items": ITEM},
            "headers": {"type": "array", "items": {"type": "string"}},
            "rows": {"type": "array", "items": {"type": "object", "properties": {
                "label": {"type": "string"}, "cells": {"type": "array", "items": {"type": "string"}}}}},
            "source": {"type": "string"}, "notes": {"type": "string"}},
            "required": ["layout", "notes"]}},
    },
    "required": ["title", "slides"],
}

SPEC_PROMPT = """You write the slide copy for a presentation for evoila US, a VMware by Broadcom specialist services firm in North America (private cloud, private AI, security and hybrid operations on VMware Cloud Foundation; assessment, build, migration and lifecycle services; delivery often under channel partners' badges). The person asking wrote an outline. Your job is to turn each outline entry into one finished slide, then call submit_deck.

WHO AND WHAT
Presenter(s): {presenters}
Audience: {audience}
Company they are presenting to (may be empty): {company}
What they want the audience to do next: {next_step}
Deck title they gave: {deck_title}
Notes about the meeting (may be empty): {notes}
Background material they pasted (may be empty): {background}

OUTLINE (one entry per slide, in order; "layout" is their pick or "auto")
{outline}

LAYOUTS (choose the one that fits each entry; honor the requester's pick unless the content cannot fit it)
- cover: first slide only. Uses the deck title and subtitle; no other fields.
- agenda: title + 3 to 6 items (heading only, under 40 characters each).
- statement: dark slide with one big message. kicker (under 30 characters, optional), statement (under 110 characters), support (one sentence under 140 characters, optional).
- bullets: title + 3 to 6 items; each item has an optional heading (2 to 4 words, bolded) and text (under 110 characters).
- columns: title + 2 or 3 items; each item has a heading (under 32 characters) and 2 to 4 points (under 75 characters each).
- cards: title + 3 or 4 items; each item has a heading (under 30 characters), text (under 95 characters) and an icon from the allowed list.
- numbers: title + 2 to 4 items; each item has figure (short, for example 40% or 3x) and heading (the label, under 60 characters); source (who published the figures). Only use figures the requester supplied. If they asked for numbers without giving them, use bracketed placeholders such as [figure] and say in the notes what to fill in.
- steps: title + 3 to 5 items; each item has a heading (under 28 characters) and text (under 90 characters).
- comparison: title + headers (2 or 3 column names, the first column holds row labels and has no header) + 2 to 5 rows; each row has a label (under 28 characters) and one short cell per header (under 60 characters).
- quote: dark slide. statement is the quote (under 160 characters, only words the requester supplied), support is the attribution.
- closing: dark slide, usually last. title (for example Next steps) + up to 3 items (heading only, under 80 characters). The presenters' contact cards are added automatically.
Every slide except cover has a title under 60 characters, written as the point of the slide, not a label.

RULES
- First slide is always cover. If the outline does not start with a cover, add one. If the requester asked for a closing or next-steps slide, use the closing layout.
- One slide per outline entry. Do not add or drop slides other than the cover.
- Write only what the requester asked for, in their direction. You may tighten and organize their words; keep their meaning.
- Never invent customers, case studies, results, figures, prices, savings, certifications, awards or dates. If the outline needs a fact you do not have, put a bracketed placeholder like [customer example] on the slide and say in the notes what is missing.
- evoila facts you may use if relevant: Broadcom Pinnacle Partner; assessment-led approach; work is delivered in phases a customer can commission one at a time; delivery often under channel partners' badges; engineering teams in multiple countries behind the US team. Do not mention VCSP or partner-of-the-year awards.
- Plain, direct, confident language. Short sentences. No hype, no buzzwords, no exclamation marks.
- No em dashes or en dashes anywhere. Use commas, colons or periods.
- notes: 2 to 4 sentences of speaker notes per slide, in the presenter's voice, saying how to land the slide. List anything in brackets that the presenter must fill in before the meeting.
- subtitle: one line under 70 characters for the cover that says what the deck is about. Do not name the company or event; the cover already says who it is prepared for.

Call submit_deck once with the whole deck."""


def _clean(t):
    t = str(t or "")
    t = t.replace("—", ", ").replace("–", "-").replace(" ,", ",")
    return re.sub(r"\s{2,}", " ", t).strip()


def _clean_obj(o):
    if isinstance(o, str):
        return _clean(o)
    if isinstance(o, list):
        return [_clean_obj(x) for x in o]
    if isinstance(o, dict):
        return {k: (_clean_obj(v) if k not in ("layout", "icon") else v) for k, v in o.items()}
    return o


def make_spec(outline, meta):
    lines = []
    for i, s in enumerate(outline.get("slides") or [], 1):
        lines.append(f"{i}. layout: {s.get('layout') or 'auto'}\n   wants: {str(s.get('content') or '').strip()}")
    prompt = SPEC_PROMPT.format(
        presenters=meta.get("presenters") or "", audience=meta.get("audience") or "",
        company=meta.get("company") or "", next_step=meta.get("next_step") or "",
        deck_title=outline.get("title") or "", notes=meta.get("notes") or "",
        background=(meta.get("background") or "")[:30000], outline="\n".join(lines) or "(empty)")
    tools = [{"name": "submit_deck", "description": "Submit the finished deck.", "input_schema": SPEC_SCHEMA}]
    messages = [{"role": "user", "content": prompt}]
    for _ in range(3):
        resp = ai_steps._call({"model": ai_steps.MODEL, "max_tokens": 12000, "tools": tools, "messages": messages},
                              timeout=240)
        content = resp.get("content", [])
        spec = None
        for b in content:
            if b.get("type") == "tool_use" and b.get("name") == "submit_deck":
                spec = b.get("input") or {}
        if spec is None:
            text = " ".join(b.get("text", "") for b in content if b.get("type") == "text")
            m = re.search(r"\{.*\}", text, re.S)
            try:
                spec = json.loads(m.group(0)) if m else None
            except Exception:
                spec = None
        if isinstance(spec, dict) and spec.get("slides"):
            return _clean_obj(spec)
        messages.append({"role": "assistant", "content": content})
        if resp.get("stop_reason") == "tool_use":
            messages.append({"role": "user", "content": [
                {"type": "tool_result", "tool_use_id": b["id"], "content": "Missing slides. Call submit_deck again with every slide."}
                for b in content if b.get("type") == "tool_use"]})
        else:
            messages.append({"role": "user", "content": "Call submit_deck now with the whole deck."})
    raise RuntimeError("Claude returned no deck")


def fallback_spec(outline):
    """No Claude (missing key or API down): lay the requester's own words out as bullet slides."""
    slides = [{"layout": "cover", "notes": "Built without the copywriting step; check every slide before the meeting."}]
    for s in outline.get("slides") or []:
        txt = str(s.get("content") or "").strip()
        parts = [p.strip(" -*•") for p in re.split(r"[\n;]+", txt) if p.strip(" -*•")]
        head = parts[0][:60] if parts else "Slide"
        rest = parts[1:] or parts
        lay = s.get("layout") if s.get("layout") in ("closing", "statement") else "bullets"
        if lay == "statement":
            slides.append({"layout": "statement", "statement": txt[:110], "notes": ""})
        elif lay == "closing":
            slides.append({"layout": "closing", "title": "Next steps", "items": [{"heading": p[:80]} for p in parts[:3]], "notes": ""})
        else:
            slides.append({"layout": "bullets", "title": head, "items": [{"text": p[:110]} for p in rest[:6]], "notes": ""})
    return _clean_obj({"title": outline.get("title") or "evoila", "subtitle": "", "slides": slides})


# ---------------------------------------------------------------- rendering

def _fit(t, n, base, low):
    l = len(t or "")
    return base if l <= n else max(low, round(base * math.sqrt(n / l), 1))


def render(spec, tpl_path, out, meta, engine_dir):
    sys.path.insert(0, engine_dir)
    import build_library as BL
    import pptkit as K
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.dml import MSO_LINE_DASH_STYLE as DASH
    VIOLET, VIOLET2, GREEN, BLUE, LBLUE, INDIGO, CREAM, WHITE = (K.VIOLET, K.VIOLET2, K.GREEN, K.BLUE, K.LBLUE,
                                                                 K.INDIGO, K.CREAM, K.WHITE)
    SOFT, MUTED = BL.SOFT, BL.MUTED
    box, text, text_in, icon, icon_disc, line, pill = K.box, K.text, K.text_in, K.icon, K.icon_disc, K.line, K.pill

    BL.TPL = Presentation(tpl_path)
    prs = Presentation(tpl_path)
    sl = prs.slides._sldIdLst
    for sid in list(sl):
        prs.part.drop_rel(sid.rId); sl.remove(sid)
    new = BL.new_slide
    X0, W = 0.68, 11.97
    ACC = [GREEN, LBLUE, INDIGO, BLUE]

    def items(s, lo=0, hi=9):
        return [i for i in (s.get("items") or []) if isinstance(i, dict)][:hi]

    def ic(it, default="CircleCheck"):
        n = it.get("icon")
        return n if n in ICON_NAMES else BL.pick_icon((it.get("heading") or "") + " " + (it.get("text") or ""), default)

    from pptx.util import Inches as _Inch

    def cover(s):
        sld = new(prs, "cover")
        for ph in sld.placeholders:
            if ph.placeholder_format.idx == 0:
                ph.text = spec.get("title") or "evoila"
            elif ph.placeholder_format.idx == 1:
                ph.text = "EVOILA  |  NORTH AMERICA"
        sub = spec.get("subtitle") or ""
        CW = 4.25                                 # clear of the cover artwork; stacked from the bottom up
        people = meta.get("people") or []
        lines_p = [[(f"{p['name']}, {p['title']}", {})] for p in people]
        ph_ = 0.27 * max(1, len(lines_p))
        py_ = 7.02 - ph_
        if lines_p:
            text(sld, 1.0, py_, CW, ph_, lines_p, size=12, color=SOFT)
        ml = [[(x, {"size": 14 if i == 0 else 12, "bold": i == 0, "color": GREEN if i == 0 else SOFT})]
              for i, x in enumerate([x for x in (meta.get("meeting_lines") or []) if x])]
        y = py_ - 0.3
        if ml:
            mh = sum(BL.para_h(m, CW, 12) for m in ml) + 0.05 * (len(ml) - 1) + 0.08
            y -= mh
            text(sld, 1.0, y, CW, mh, ml, size=13, color=GREEN, spacing=3)
            y -= 0.22
        if sub:
            sz = 15 if len(sub) <= 44 else 14
            sh_ = BL.para_h(sub, CW + 0.4, sz) + 0.05
            y -= sh_
            text(sld, 1.0, y, CW + 0.4, sh_, sub, size=sz, color=SOFT)
        ttl = spec.get("title") or "evoila"
        ts = 40
        while True:
            th = BL.n_lines(ttl, 5.35, ts) * ts * 0.95 / 72 + 0.08
            if y - 0.2 - th - 0.5 >= 1.75 or ts <= 26:
                break
            ts -= 2
        ty = y - 0.2 - th
        for ph in sld.placeholders:
            if ph.placeholder_format.idx == 0:
                ph.left, ph.top, ph.width, ph.height = _Inch(1.0), _Inch(ty), _Inch(5.35), _Inch(th)
                for r in ph.text_frame.paragraphs[0].runs:
                    r.font.size = Pt(ts)
            elif ph.placeholder_format.idx == 1:
                ph.left, ph.top, ph.width, ph.height = _Inch(1.0), _Inch(ty - 0.42), _Inch(4.4), _Inch(0.24)
        return sld

    def agenda(s):
        sld = new(prs, "light", s.get("title") or "Agenda")
        its = items(s, hi=7)
        h = min(0.72, 4.9 / max(1, len(its)))
        for i, it in enumerate(its):
            y = 1.55 + i * (h + 0.08)
            box(sld, X0, y, W, h, fill=CREAM, radius=0.18)
            text(sld, X0 + 0.25, y, 0.8, h, f"{i + 1:02d}", size=20, bold=True, color=BLUE, anchor=MSO_ANCHOR.MIDDLE)
            text(sld, X0 + 1.2, y, W - 1.5, h, it.get("heading") or it.get("text") or "", size=18, bold=True,
                 color=VIOLET, anchor=MSO_ANCHOR.MIDDLE)
        return sld

    def statement(s):
        st = s.get("statement") or s.get("title") or ""
        ttl = s.get("title") or ""
        if ttl.strip(" .").lower() in st.strip(" .").lower() or st.strip(" .").lower().startswith(ttl.strip(" .").lower()[:30]):
            ttl = ""                             # don't print the same line twice
        sld = new(prs, "dark", ttl)
        BL.add_blob(sld, "right")
        y = 2.0
        if s.get("kicker"):
            text(sld, X0, y, 7.6, 0.4, s["kicker"].upper(), size=13, bold=True, color=GREEN)
            y += 0.55
        text(sld, X0, y, 7.6, 2.4, st, size=_fit(st, 70, 36, 24), bold=True, color=WHITE)
        if s.get("support"):
            text(sld, X0, 5.0, 7.4, 1.0, s["support"], size=17, color=SOFT)
        return sld

    def bullets(s):
        sld = new(prs, "light", s.get("title") or "")
        y0 = 1.5
        if s.get("support"):
            text(sld, X0, y0, W, 0.5, s["support"], size=15, color=MUTED)
            y0 += 0.65
        its = items(s, hi=6)
        gap = min(0.95, (6.55 - y0) / max(1, len(its)))
        for i, it in enumerate(its):
            y = y0 + i * gap
            box(sld, X0, y, W, gap - 0.12, fill=CREAM, radius=0.18)
            box(sld, X0 + 0.25, y + (gap - 0.12) / 2 - 0.08, 0.16, 0.16, fill=ACC[i % 2], shape=MSO_SHAPE.OVAL)
            runs = []
            if it.get("heading"):
                runs.append((it["heading"].rstrip(":") + "   ", {"bold": True, "color": VIOLET}))
            runs.append((it.get("text") or "", {"color": VIOLET}))
            text(sld, X0 + 0.65, y, W - 0.9, gap - 0.12, [runs], size=_fit(it.get("text") or "", 90, 19, 14),
                 anchor=MSO_ANCHOR.MIDDLE)
        return sld

    def columns(s):
        sld = new(prs, "light", s.get("title") or "")
        its = items(s, hi=3) or [{}]
        n = len(its); g = 0.3; w = (W - g * (n - 1)) / n
        for i, it in enumerate(its):
            x = X0 + i * (w + g)
            box(sld, x, 1.55, w, 4.95, fill=CREAM, radius=0.06)
            box(sld, x, 1.55, w, 0.09, fill=ACC[i % len(ACC)], radius=0)
            text(sld, x + 0.3, 1.85, w - 0.6, 0.8, it.get("heading") or "", size=20, bold=True, color=VIOLET)
            pts = it.get("points") or ([it["text"]] if it.get("text") else [])
            paras = [[("•  ", {"color": BLUE, "bold": True}), (p, {})] for p in pts[:4]]
            if paras:
                text(sld, x + 0.3, 2.75, w - 0.6, 3.5, paras, size=_fit(max(pts, key=len), 60, 15, 12),
                     color=VIOLET, spacing=10)
        return sld

    def cards(s):
        sld = new(prs, "light", s.get("title") or "")
        its = items(s, hi=4) or [{}]
        n = len(its); g = 0.3; w = (W - g * (n - 1)) / n
        for i, it in enumerate(its):
            x = X0 + i * (w + g)
            box(sld, x, 1.65, w, 4.6, fill=CREAM, radius=0.08)
            col = ACC[i % len(ACC)]
            icon_disc(sld, ic(it), x + 0.35, 1.95, 0.95, disc=col, color="v" if col in (GREEN, LBLUE) else "w")
            text(sld, x + 0.35, 3.15, w - 0.7, 0.8, it.get("heading") or "", size=_fit(it.get("heading") or "", 22, 19, 15),
                 bold=True, color=VIOLET)
            text(sld, x + 0.35, 4.0, w - 0.7, 2.1, it.get("text") or "", size=_fit(it.get("text") or "", 80, 15, 12),
                 color=MUTED)
        return sld

    def numbers(s):
        sld = new(prs, "dark", s.get("title") or "")
        its = items(s, hi=4) or [{}]
        n = len(its); g = 0.35; w = (W - g * (n - 1)) / n
        for i, it in enumerate(its):
            x = X0 + i * (w + g)
            box(sld, x, 1.8, w, 3.7, fill=VIOLET2, alpha=92, line="FFFFFF@20", lw=0.75, radius=0.08)
            fig = it.get("figure") or "[figure]"
            text(sld, x + 0.3, 2.1, w - 0.6, 1.3, fig, size=_fit(fig, 6, 54, 30), bold=True, color=GREEN)
            text(sld, x + 0.3, 3.55, w - 0.6, 1.8, it.get("heading") or it.get("text") or "",
                 size=_fit(it.get("heading") or "", 55, 17, 13), color=WHITE)
        if s.get("source"):
            text(sld, X0, 5.85, W, 0.4, "Source: " + s["source"], size=11, color=SOFT)
        return sld

    def steps(s):
        sld = new(prs, "light", s.get("title") or "")
        its = items(s, hi=5) or [{}]
        n = len(its); g = 0.25; w = (W - g * (n - 1)) / n
        for i, it in enumerate(its):
            x = X0 + i * (w + g)
            d = box(sld, x, 1.75, 0.95, 0.95, grad=(GREEN, BLUE, 45), shape=MSO_SHAPE.OVAL)
            text_in(d, str(i + 1), size=26, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)
            if i < n - 1:
                line(sld, x + 1.05, 2.22, x + w + g - 0.1, 2.22, color="1F1534@35", width=1.5, dash=DASH.DASH)
            text(sld, x, 3.0, w - 0.1, 0.8, it.get("heading") or "", size=_fit(it.get("heading") or "", 22, 19, 15),
                 bold=True, color=VIOLET)
            text(sld, x, 3.85, w - 0.1, 2.4, it.get("text") or "", size=_fit(it.get("text") or "", 75, 15, 12), color=MUTED)
        return sld

    def comparison(s):
        sld = new(prs, "light", s.get("title") or "")
        heads = (s.get("headers") or [])[:3]
        rows = [r for r in (s.get("rows") or []) if isinstance(r, dict)][:5]
        nc = max(1, len(heads)); lw = 2.8; g = 0.12
        cw = (W - lw - g * nc) / nc
        rh = min(0.95, (6.5 - 2.3) / max(1, len(rows)))
        for j, h in enumerate(heads):
            x = X0 + lw + g + j * (cw + g)
            b = box(sld, x, 1.55, cw, 0.62, fill=[INDIGO, BLUE, GREEN][j % 3], radius=0.12)
            text_in(b, h, size=15, bold=True, color=WHITE if j < 2 else VIOLET, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        for i, r in enumerate(rows):
            y = 2.3 + i * (rh + g / 2)
            b = box(sld, X0, y, lw, rh, fill=VIOLET, radius=0.08)
            text_in(b, r.get("label") or "", size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
            for j in range(nc):
                cell = (r.get("cells") or [])[j] if j < len(r.get("cells") or []) else ""
                x = X0 + lw + g + j * (cw + g)
                b = box(sld, x, y, cw, rh, fill=CREAM, radius=0.08)
                text_in(b, cell, size=_fit(cell, 50, 13.5, 11), color=VIOLET, anchor=MSO_ANCHOR.MIDDLE)
        return sld

    def quote(s):
        st = s.get("statement") or s.get("title") or ""
        ttl = s.get("title") or ""
        if ttl.strip(" .").lower() in st.strip(" .").lower() or st.strip(" .").lower().startswith(ttl.strip(" .").lower()[:30]):
            ttl = ""                             # don't print the same line twice
        sld = new(prs, "dark", ttl)
        BL.add_blob(sld, "right")
        icon(sld, "Quote", X0, 1.7, 0.8, "g")
        q = s.get("statement") or ""
        text(sld, X0, 2.65, 7.8, 2.4, q, size=_fit(q, 110, 28, 20), bold=True, color=WHITE)
        if s.get("support"):
            text(sld, X0, 5.25, 7.6, 0.5, s["support"], size=16, color=GREEN)
        return sld

    def closing(s):
        sld = new(prs, "dark", s.get("title") or "Next steps")
        its = items(s, hi=3)
        gap = 4.5 / max(3, len(its) or 1)
        for i, it in enumerate(its):
            y = 1.55 + i * gap
            d = box(sld, X0, y, 0.9, 0.9, grad=(GREEN, BLUE, 45), shape=MSO_SHAPE.OVAL)
            text_in(d, str(i + 1), size=26, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)
            if i < len(its) - 1:
                line(sld, X0 + 0.45, y + 0.95, X0 + 0.45, y + gap - 0.05, color="FFFFFF@40", width=1.5, dash=DASH.DASH)
            text(sld, X0 + 1.17, y, 4.8, 0.9, it.get("heading") or it.get("text") or "", size=17, bold=True, color=WHITE,
                 anchor=MSO_ANCHOR.MIDDLE)
        for i, c in enumerate((meta.get("people") or [])[:2]):
            y = 1.55 + i * 2.05
            box(sld, 6.95, y, 5.7, 1.8, fill=VIOLET2, alpha=92, line="FFFFFF@20", lw=0.75, radius=0.08)
            box(sld, 7.2, y + 0.3, 1.2, 1.2, fill="FFFFFF@12", shape=MSO_SHAPE.OVAL)
            icon(sld, "UserCheck", 7.5, y + 0.6, 0.6, "g")
            text(sld, 8.65, y + 0.3, 3.8, 0.45, c["name"], size=16, bold=True, color=WHITE)
            text(sld, 8.65, y + 0.78, 3.8, 0.35, c["title"], size=13, color=SOFT)
            text(sld, 8.65, y + 1.15, 3.8, 0.35, c.get("email") or "", size=13, color=GREEN)
        pill(sld, 6.95, 5.85, 2.6, 0.5, "evoila.com/us", fill=GREEN, color=VIOLET, size=14)
        return sld

    # ---- layouts the deck converter uses (convert_deck.py); the outline builder can use them too
    from pptx.util import Inches as _In

    def table(s):
        sld = new(prs, "light", s.get("title") or "")
        heads = [str(h) for h in (s.get("headers") or [])][:6]
        rows = [r for r in (s.get("rows") or []) if isinstance(r, dict)][:8]
        nc = max([len(heads)] + [1 + len(r.get("cells") or []) for r in rows])
        nc = max(2, min(6, nc))
        heads = (heads + [""] * nc)[:nc]
        labelled = any((r.get("label") or "").strip() for r in rows)
        g = 0.08
        lw = min(3.0, W / nc * 1.15) if labelled else (W - g * (nc - 1)) / nc
        cw = (W - lw - g * (nc - 1)) / max(1, nc - 1)
        widths = [lw] + [cw] * (nc - 1)
        xs = [X0]
        for w_ in widths[:-1]:
            xs.append(xs[-1] + w_ + g)
        top = 1.55
        hh = 0.62 if any(heads) else 0
        rh = min(0.82, (6.55 - top - hh - g) / max(1, len(rows)) - g)
        longest = max([len(str(c)) for r in rows for c in ([r.get("label") or ""] + list(r.get("cells") or []))] + [8])
        fs = _fit("x" * longest, 28, 15, 10.5)
        if hh:
            for j, h in enumerate(heads):
                if not h and j == 0 and labelled:
                    continue
                b = box(sld, xs[j], top, widths[j], hh, fill=INDIGO, radius=0.1)
                text_in(b, h, size=_fit(h, 22, 14, 11), bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                        anchor=MSO_ANCHOR.MIDDLE, margin=0.08)
        for i, r in enumerate(rows):
            y = top + hh + g + i * (rh + g)
            cells = [r.get("label") or ""] + list(r.get("cells") or []) if labelled else list(r.get("cells") or [])
            cells = (cells + [""] * nc)[:nc]
            for j, c in enumerate(cells):
                first = labelled and j == 0
                b = box(sld, xs[j], y, widths[j], rh, fill=VIOLET if first else CREAM, radius=0.06)
                text_in(b, str(c), size=fs, bold=first, color=WHITE if first else VIOLET, anchor=MSO_ANCHOR.MIDDLE,
                        align=PP_ALIGN.LEFT if first else PP_ALIGN.CENTER, margin=0.1)
        if s.get("source"):
            text(sld, X0, 6.62, W, 0.3, "Source: " + s["source"], size=10.5, color=MUTED)
        return sld

    def chart(s):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        ch = s.get("chart") or {}
        cats = [str(c) for c in ch.get("categories") or []]
        series = [x for x in ch.get("series") or [] if x.get("values")][:6]
        if not cats or not series:
            return bullets(dict(s, items=[{"text": "[chart data]"}]))
        kind = s.get("chart_kind") or ""
        orig = (ch.get("type") or "").upper()
        if not kind:
            kind = "line" if "LINE" in orig else "pie" if ("PIE" in orig or "DOUGHNUT" in orig) else "bar" if orig.startswith("BAR") else "column"
        dark = False
        sld = new(prs, "light", s.get("title") or "")
        cd = CategoryChartData()
        cd.categories = cats
        for x in (series[:1] if kind == "pie" else series):
            cd.add_series(x.get("name") or "", [float(v or 0) for v in x["values"][:len(cats)]])
        xt = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
              "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.DOUGHNUT}[kind]
        side = bool(s.get("support"))
        cw_ = 8.0 if side else W
        gf = sld.shapes.add_chart(xt, _In(X0), _In(1.5), _In(cw_), _In(4.95), cd)
        c = gf.chart
        c.font.name = "Inter"
        c.font.size = Pt(12)
        c.font.color.rgb = K.rgb(VIOLET)
        pal = [BLUE, GREEN, INDIGO, LBLUE, "8C84A0", "1F1534"]
        multi = len(series) > 1 and kind != "pie"
        c.has_legend = multi or kind == "pie"
        if c.has_legend:
            c.legend.position = XL_LEGEND_POSITION.BOTTOM if kind != "pie" else XL_LEGEND_POSITION.RIGHT
            c.legend.include_in_layout = False
            c.legend.font.size = Pt(12)
        plot = c.plots[0]
        if kind == "pie":
            for k, pt in enumerate(plot.series[0].points):
                pt.format.fill.solid(); pt.format.fill.fore_color.rgb = K.rgb(pal[k % len(pal)])
            plot.has_data_labels = True
            plot.data_labels.font.size = Pt(12); plot.data_labels.font.color.rgb = K.rgb(WHITE)
        else:
            for k, se in enumerate(plot.series):
                col = K.rgb(pal[k % len(pal)])
                if kind == "line":
                    se.format.line.color.rgb = col; se.format.line.width = Pt(3)
                    se.smooth = False
                    try:
                        se.marker.format.fill.solid(); se.marker.format.fill.fore_color.rgb = col; se.marker.format.line.color.rgb = col
                    except Exception:
                        pass
                else:
                    se.format.fill.solid(); se.format.fill.fore_color.rgb = col
            if kind in ("column", "bar"):
                try:
                    plot.gap_width = 60; plot.overlap = -10 if multi else 0
                except Exception:
                    pass
            if len(cats) * len(series) <= 12:
                plot.has_data_labels = True
                dl = plot.data_labels
                dl.font.size = Pt(11); dl.font.color.rgb = K.rgb(VIOLET)
                try:
                    dl.number_format = '#,##0.##'; dl.number_format_is_linked = False
                except Exception:
                    pass
            va = c.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = K.rgb("E4DDD6")
            va.format.line.fill.background()
            va.tick_labels.font.size = Pt(11); va.tick_labels.font.color.rgb = K.rgb(MUTED)
            ca = c.category_axis
            ca.format.line.color.rgb = K.rgb("CFC6BD")
            ca.tick_labels.font.size = Pt(12)
        if side:
            b = box(sld, X0 + cw_ + 0.35, 1.5, W - cw_ - 0.35, 4.95, fill=VIOLET, radius=0.08)
            text(sld, X0 + cw_ + 0.65, 1.85, W - cw_ - 0.95, 0.4, "WHAT IT SHOWS", size=12, bold=True, color=GREEN)
            text(sld, X0 + cw_ + 0.65, 2.35, W - cw_ - 0.95, 3.8, s["support"], size=_fit(s["support"], 80, 21, 16), bold=True, color=WHITE)
        if s.get("source"):
            text(sld, X0, 6.55, W, 0.3, "Source: " + s["source"], size=10.5, color=MUTED)
        return sld

    def image(s):
        from PIL import Image as _Im
        sld = new(prs, "light", s.get("title") or "")
        pth = s.get("image_path")
        its = items(s, hi=4)
        cap = s.get("caption") or ""
        if its:
            fx, fy, fw, fh = X0, 1.5, 7.55, 5.0
        else:
            fx, fy, fw, fh = X0, 1.5, W, 4.75 if cap else 5.05
        frame = box(sld, fx, fy, fw, fh, fill=CREAM, radius=0.06)
        if pth and os.path.exists(pth):
            iw, ih = _Im.open(pth).size
            pad = 0.22
            aw, ah = fw - 2 * pad, fh - 2 * pad
            sc_ = min(aw / iw, ah / ih)
            w_, h_ = iw * sc_, ih * sc_
            pic = sld.shapes.add_picture(pth, _In(fx + (fw - w_) / 2), _In(fy + (fh - h_) / 2), _In(w_), _In(h_))
            pic.line.color.rgb = K.rgb("E4DDD6"); pic.line.width = Pt(0.75)
        if cap and not its:
            text(sld, X0, fy + fh + 0.12, W, 0.4, cap, size=13, color=MUTED)
        if its:
            tx = fx + fw + 0.35; tw = X0 + W - tx
            if cap:
                text(sld, tx, 1.5, tw, 0.9, cap, size=_fit(cap, 70, 17, 14), bold=True, color=VIOLET)
            y = 1.5 + (1.05 if cap else 0)
            gap = min(1.15, (6.5 - y) / max(1, len(its)))
            for i, it in enumerate(its):
                box(sld, tx, y + i * gap + 0.12, 0.14, 0.14, fill=ACC[i % 2], shape=MSO_SHAPE.OVAL)
                runs = []
                if it.get("heading"):
                    runs.append((it["heading"].rstrip(":") + "  ", {"bold": True}))
                runs.append((it.get("text") or "", {}))
                text(sld, tx + 0.35, y + i * gap, tw - 0.35, gap - 0.08, [runs], size=_fit(it.get("text") or "", 70, 15, 12), color=VIOLET)
        return sld

    def hub(s):
        sld = new(prs, "light", s.get("title") or "")
        its = items(s, hi=6) or [{}]
        cx, cy, rx, ry = 6.67, 4.05, 4.25, 1.85
        n = len(its)
        import math as _m
        pos = []
        for i in range(n):
            a = -_m.pi / 2 + 2 * _m.pi * i / n
            pos.append((cx + rx * _m.cos(a), cy + ry * _m.sin(a)))
        for (px, py) in pos:
            line(sld, cx, cy, px, py, color="0075E4@55", width=1.25, dash=DASH.DASH)
        d = box(sld, cx - 1.05, cy - 1.05, 2.1, 2.1, grad=(GREEN, BLUE, 45), shape=MSO_SHAPE.OVAL)
        hc = s.get("hub_center") or ""
        text_in(d, hc, size=_fit(hc, 16, 18, 13), bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.15)
        bw, bh = 2.75, 0.95
        for i, ((px, py), it) in enumerate(zip(pos, its)):
            bx, by = min(max(X0, px - bw / 2), X0 + W - bw), min(max(1.45, py - bh / 2), 6.55 - bh)
            box(sld, bx, by, bw, bh, fill=WHITE, line=VIOLET, lw=1.0, radius=0.14)
            col = ACC[i % len(ACC)]
            icon_disc(sld, ic(it), bx + 0.14, by + 0.17, 0.6, disc=col, color="v" if col in (GREEN, LBLUE) else "w")
            hd = it.get("heading") or it.get("text") or ""
            text(sld, bx + 0.86, by, bw - 0.95, bh, hd, size=_fit(hd, 20, 14, 11), bold=True, color=VIOLET, anchor=MSO_ANCHOR.MIDDLE)
        return sld

    draw = dict(cover=cover, agenda=agenda, statement=statement, bullets=bullets, columns=columns, cards=cards,
                numbers=numbers, steps=steps, comparison=comparison, quote=quote, closing=closing,
                table=table, chart=chart, image=image, hub=hub)
    slides = [s for s in spec.get("slides") or [] if isinstance(s, dict)]
    if not slides or slides[0].get("layout") != "cover":
        slides = [{"layout": "cover", "notes": ""}] + slides
    for s in slides:
        sld = draw.get(s.get("layout"), bullets)(s)
        if s.get("notes"):
            K.notes(sld, s["notes"])
    BL.set_sections(prs, [("Deck", [e.get("id") for e in prs.slides._sldIdLst])])
    prs.save(out)
    return len(slides)


# ---------------------------------------------------------------- build flow

SOURCE_KINDS = {
    "Analyst firms": ("gartner.com idc.com forrester.com mckinsey.com bcg.com bain.com deloitte.com pwc.com kpmg.com "
                      "ey.com accenture.com").split(),
    "News and trade press": ("hipaajournal.com healthcareitnews.com fiercehealthcare.com beckershospitalreview.com "
                             "modernhealthcare.com statnews.com healthcaredive.com cybersecuritydive.com ciodive.com "
                             "bankingdive.com utilitydive.com reuters.com apnews.com bloomberg.com wsj.com ft.com nytimes.com "
                             "washingtonpost.com cnbc.com axios.com forbes.com fortune.com darkreading.com securityweek.com "
                             "theregister.com bleepingcomputer.com scmagazine.com infosecurity-magazine.com techtarget.com "
                             "computerweekly.com computerworld.com networkworld.com cio.com csoonline.com infoworld.com zdnet.com "
                             "crn.com govtech.com statescoop.com nextgov.com fedscoop.com route-fifty.com govexec.com "
                             "americanbanker.com").split(),
}


def _kind(url):
    import urllib.parse as up
    h = up.urlparse(url).netloc.lower().split(":")[0]
    h = h[4:] if h.startswith("www.") else h
    if re.search(r"\.(gov|mil|int)(\.[a-z]{2})?$", h):
        return "Government and regulators"
    for k, doms in SOURCE_KINDS.items():
        if any(h == d or h.endswith("." + d) for d in doms):
            return k
    return "Industry studies and associations"


def research_slide(entry, meta, report):
    """Research one outline entry and return a numbers slide (verified figures or placeholders)."""
    rules = entry.get("research") or {}
    try:
        months = max(3, min(36, int(rules.get("months") or 12)))
    except Exception:
        months = 12
    rules["months"] = months
    topic = str(entry.get("content") or "").strip()
    today = datetime.date.today()
    try:
        r = ai_steps.make_topic_research(topic, rules, meta, today.isoformat())
    except Exception as e:
        r = {"stats": []}
        report.append(f"Research slide '{topic[:60]}': research failed ({str(e)[:150]}).")
    allowed = set(rules.get("sources") or [])
    used, seen = [], set()
    for st in (r.get("stats") or [])[:8]:
        label = f"{st.get('figure','')} {str(st.get('claim',''))[:70]} ({st.get('source','')})"
        if allowed and _kind(st.get("url", "")) not in allowed:
            report.append(f"DROPPED (source kind not allowed by the requester): {label} {st.get('url','')}")
            continue
        key = (rr.compact_figure(st.get("figure", "")), str(st.get("source", "")).lower())
        if key in seen or len(used) >= 4:
            continue
        ok, why = rr.check_stat(st, today, max_days=int(months * 30.5))
        if not ok:
            report.append(f"DROPPED ({why}): {label} {st.get('url','')}")
            continue
        seen.add(key)
        used.append(st)
        report.append(f"USED (found on the page, {st.get('year','')}): {label} {st.get('url','')}")
    title = rr.nodash(r.get("title") or topic[:60] or "What the data shows")
    if len(used) >= 2:
        srcs = []
        for st in used:
            s_ = f"{st.get('source','')}, {st.get('year','')}".strip(", ")
            if s_ not in srcs:
                srcs.append(s_)
        notes = " ".join(x for x in [rr.nodash(r.get("talking_point") or ""), rr.nodash(r.get("takeaway") or "")] if x)
        notes += " Sources: " + "; ".join(f"{st.get('source','')} ({st.get('url','')})" for st in used)
        report.append(f"Research slide '{title}': included with {len(used)} confirmed figures.")
        return {"layout": "numbers", "title": title, "source": "; ".join(srcs)[:220],
                "items": [{"figure": st.get("figure", ""), "heading": st.get("claim", "")} for st in used], "notes": notes}
    report.append(f"Research slide '{title}': fewer than two figures passed the checks, placeholders left on the slide.")
    return {"layout": "numbers", "title": title, "source": "[source]",
            "items": [{"figure": "[figure]", "heading": "[what the figure shows]"} for _ in range(3)],
            "notes": "We could not confirm enough recent figures that met your rules for this slide, so it has placeholders. "
                     "The build report on the request lists what was found and why each figure was dropped."}


def run():
    os.makedirs(rr.WORK, exist_ok=True)
    f = rr.get_record(rr.T_REQ, rr.REC)
    try:
        outline = json.loads(f.get(F_OUTLINE) or "{}")
    except Exception:
        raise RuntimeError("The custom outline on this row is not valid JSON.")
    people = []
    for pid in f.get(rr.F["presenters"], [])[:2]:
        p = rr.get_record(rr.T_PRES, pid)
        people.append({"name": p.get(rr.P["name"], ""), "title": p.get(rr.P["title"], ""), "email": p.get(rr.P["email"], "")})
    prospect = (f.get(rr.F["prospect"]) or "").strip()
    when = ""
    if f.get(rr.F["date"]):
        d = datetime.date.fromisoformat(f[rr.F["date"]])
        when = f"{d.strftime('%B')} {d.day}, {d.year}"
    meta = {"people": people, "presenters": ", ".join(f"{p['name']} ({p['title']})" for p in people),
            "audience": f.get(rr.F["audience"]) or "", "company": prospect,
            "next_step": f.get(F_NEXT) or "", "notes": f.get(rr.F["notes"]) or "",
            "background": f.get(F_BACKGROUND) or "",
            "meeting_lines": [f"Prepared for {prospect}" if prospect else "", when]}
    note = ""
    meta["deck_title"] = outline.get("title") or ""
    entries = outline.get("slides") or []
    research_idx = [i for i, e in enumerate(entries) if e.get("layout") == "research"]
    report, researched = [], {}
    if research_idx:
        import concurrent.futures as cf
        with cf.ThreadPoolExecutor(max_workers=3) as ex:
            futs = {i: ex.submit(research_slide, entries[i], meta, report) for i in research_idx}
            for i, fu in futs.items():
                researched[i] = fu.result()
        # the copywriter sees a plain numbers slide it must leave alone; the researched slide replaces it afterwards
        outline = dict(outline, slides=[dict(e, layout="numbers", content="[Research slide, filled in by the build. Write a one-line title only.]")
                                        if i in researched else e for i, e in enumerate(entries)])
    try:
        spec = make_spec(outline, meta)
    except Exception as e:
        spec = fallback_spec(outline)
        note = f"Claude copywriting failed, slides use the outline text as written: {str(e)[:300]}"
    if researched:
        offset = len(spec.get("slides", [])) - len(entries)
        for i, sl in researched.items():
            j = i + max(offset, 0)
            if 0 <= j < len(spec.get("slides", [])):
                spec["slides"][j] = sl
        note = "\n".join(([note] if note else []) + report)
    rr.update({rr.F["plan"]: json.dumps(spec, indent=1)[:95000]})
    tpl = rr.fetch_assets()
    stamp = f.get(rr.F["date"]) or datetime.date.today().isoformat()
    base = f"evoila_{rr.safe(outline.get('title') or prospect or 'Deck')}_{stamp}"
    pptx = os.path.join(rr.WORK, base + ".pptx")
    n = render(spec, tpl, pptx, meta, rr.ENGINE)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", rr.WORK, pptx],
                   check=True, timeout=240, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf = pptx[:-5] + ".pdf"
    rr.update({rr.F["pptx"]: [], rr.F["pdf"]: []})
    rr.upload(rr.F["pptx"], pptx, "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    rr.upload(rr.F["pdf"], pdf, "application/pdf")
    rr.update({rr.F["status"]: "Ready", rr.F["log"]: rr.RUN_URL or None, rr.F["error"]: "",
               rr.F["libver"]: "custom outline", rr.F["rcheck"]: note or None})
    print("custom deck done:", n, "slides")
